import sys
from pathlib import Path
from docx import Document

# Check exp_log.md
log_path = Path('/Users/xinyutan/Documents/量化投资/quant-project/experiments/exp_log.md')
with open(log_path, 'r', encoding='utf-8') as f:
    content = f.read()

old_patterns = ['1147.79%', '19.23%', '21.36%', '7.49%', '14.13%', '17.21%']
print('=== exp_log.md 旧数据检查 ===')
for p in old_patterns:
    count = content.count(p)
    if count > 0:
        print(f'  发现 {p}: {count} 次')
    else:
        print(f'  未发现 {p}')

# Check docx
doc_path = Path('/Users/xinyutan/Documents/量化投资/quant-project/results/策略总结与优化历程.docx')
doc = Document(doc_path)
full_text = '\n'.join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        full_text += '\n' + ' | '.join([c.text for c in row.cells])

print('\n=== docx 旧数据检查 ===')
for p in old_patterns:
    count = full_text.count(p)
    if count > 0:
        print(f'  发现 {p}: {count} 次')
    else:
        print(f'  未发现 {p}')

# Check new data
new_patterns = ['655.95%', '19.88%', '20.65%', '12.27%', '17.08%', '0.72']
print('\n=== 新数据存在性检查 ===')
all_ok = True
for p in new_patterns:
    in_md = p in content
    in_docx = p in full_text
    status = 'OK' if (in_md or in_docx) else 'MISSING'
    if not (in_md or in_docx):
        all_ok = False
    print(f'  {p}: exp_log={in_md}, docx={in_docx} [{status}]')

if all_ok:
    print('\n所有验证通过！')
else:
    print('\n部分数据缺失，请检查。')
    sys.exit(1)
