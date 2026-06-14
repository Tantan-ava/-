"""
Append lookahead bias fix section to the docx report
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

doc_path = '/Users/xinyutan/Documents/量化投资/quant-project/results/策略总结与优化历程.docx'
doc = Document(doc_path)

# Add a page break before new section
doc.add_page_break()

# Title
heading = doc.add_heading('前视偏差系统性审查与修复', level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Date
date_p = doc.add_paragraph()
date_p.add_run(f'记录日期: {datetime.now().strftime("%Y-%m-%d")}').italic = True

# Background
doc.add_heading('背景', level=2)
doc.add_paragraph(
    '在策略代码审查中发现多处潜在的前视偏差（Look-ahead Bias），即使用未来信息影响当前决策。 '
    '这些问题分布在滚动窗口计算、扩展窗口阈值校准、ML信号对齐、风险平价权重和流动性过滤等多个模块。 '
    '本次审查对代码库进行了系统性修复，并通过自动化脚本验证修复逻辑的正确性。'
)

# Fix list
doc.add_heading('修复文件清单', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '文件路径'
hdr_cells[1].text = '修复内容'
hdr_cells[2].text = '修复类型'

fixes = [
    ('src/run_optimization_experiments.py',
     '滚动夏普均值/标准差添加 .shift(1)；组合波动率添加 .shift(1)；风险预算方差添加 .shift(1)；风险平价历史窗口改为 iloc[start_idx:date_idx] 排除当前日期',
     '计算方法'),
    ('src/run_week8_tutorial.py',
     '市值代理因子 mcap_proxy 的 rolling(12).std() 添加 .shift(1)',
     '计算方法'),
    ('src/run_week7_experiments.py',
     'ML信号 ml_signal.fillna(0) 添加 .shift(1) 对齐持仓周期',
     '逻辑对齐'),
    ('src/run_week7_ml_and_combo.py',
     'ML信号添加 .shift(1) 对齐持仓周期',
     '逻辑对齐'),
    ('src/run_experiments.py',
     '流动性过滤改为使用初始24个月数据计算统计量，避免全样本前视偏差',
     '静态筛选'),
    ('src/strategy/daily_tactician.py',
     'Z-score的 expanding().mean/std 添加 .shift(1)；阈值校准的 expanding().quantile() 添加 .shift(1)；增加NaN阈值边界保护',
     '扩展窗口'),
    ('src/run_epu_weight_experiment.py',
     'EPU扩展窗口统计量（median/mean/p75/p25）及min/max归一化均添加 .shift(1)',
     '扩展窗口'),
    ('src/run_turnover_epu_experiment.py',
     'EPU阈值及中位数权重调整的 expanding() 统计添加 .shift(1)',
     '扩展窗口'),
]

for filepath, content, ftype in fixes:
    row_cells = table.add_row().cells
    row_cells[0].text = filepath
    row_cells[1].text = content
    row_cells[2].text = ftype

# Validation
doc.add_heading('验证结果', level=2)
doc.add_paragraph(
    '编写 validate_lookahead_fixes.py 对7个核心修改点进行自动化验证，结果如下：'
)

v_table = doc.add_table(rows=1, cols=2)
v_table.style = 'Light Grid Accent 1'
v_hdr = v_table.rows[0].cells
v_hdr[0].text = '验证项'
v_hdr[1].text = '状态'

validations = [
    ('风险平价历史窗口不包含当前日期', '通过'),
    ('滚动波动率/均值已添加 shift(1)', '通过'),
    ('扩展窗口分位数已添加 shift(1)', '通过'),
    ('ML信号预测周期已对齐持仓周期', '通过'),
    ('流动性过滤使用历史窗口而非全样本', '通过'),
    ('情绪Z-score扩展窗口已排除当日', '通过'),
    ('市值代理因子滚动标准差已添加 shift(1)', '通过'),
]

for item, status in validations:
    row = v_table.add_row().cells
    row[0].text = item
    row[1].text = status

doc.add_paragraph('结论：所有前视偏差修复逻辑验证通过。').bold = True

# Key risks
doc.add_heading('关键风险点说明', level=2)
risks = [
    ('.expanding().quantile() 的前视偏差',
     'expanding() 在日期 t 的结果包含 t 自身，若直接用于 t 的决策阈值，则使用了当日信息。修复方式：.expanding().quantile(...).shift(1)。'),
    ('风险平价权重 loc[:date] 的隐患',
     'Pandas loc[:date] 包含 date 本身，iloc[-60:] 会取到当前期。修复方式：使用索引切片 iloc[start_idx:date_idx]。'),
    ('ML预测 horizon 不匹配',
     'target = df_ret.shift(-1) 表示预测下月收益，但回测直接使用当月信号决定当月持仓。修复方式：ml_signal.shift(1)。'),
    ('静态筛选的全样本泄露',
     '流动性过滤若在回测开始前用全样本计算阈值，则后续每期持仓都隐含使用了未来信息。修复方式：仅使用回测初期的历史窗口计算。'),
]

for title, desc in risks:
    p = doc.add_paragraph()
    p.add_run(f'{title}: ').bold = True
    p.add_run(desc)

# Checklist
doc.add_heading('后续开发检查清单', level=2)
checklist = [
    '任何 rolling(X).mean/std/var/sum() 用于当期决策 → 确认是否已 .shift(1)',
    '任何 expanding().quantile/median/mean/std() 用于当期阈值 → 确认是否已 .shift(1)',
    '任何 .quantile()/.median()/.std() 直接在完整 Series 上调用 → 确认是否为纯分析/作图',
    'ML/预测模型 → 确认预测目标 y 的领先期数与回测权重应用期数匹配',
    '静态股票池筛选 → 确认筛选指标仅使用筛选时点之前的历史数据',
    '风险平价/最小方差权重 → 确认波动率计算窗口不包含当期收益',
]
for item in checklist:
    doc.add_paragraph(item, style='List Bullet')

# Save
output_path = doc_path
doc.save(output_path)
print(f"文档已更新并保存: {output_path}")
