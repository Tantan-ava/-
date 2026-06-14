from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from pathlib import Path

doc_path = Path('/Users/xinyutan/Documents/量化投资/quant-project/results/策略总结与优化历程.docx')
doc = Document(doc_path)

# Add page break
doc.add_page_break()

# Title
heading = doc.add_heading('附录：前视偏差修复与实验数据更新', level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Date
date_p = doc.add_paragraph()
date_p.add_run('更新日期: 2026-06-14').italic = True

# Background
doc.add_heading('修复背景', level=2)
doc.add_paragraph(
    '在策略代码系统性审查中发现多处潜在的前视偏差（Look-ahead Bias），即使用未来信息影响当前决策。 '
    '本次更新对代码库进行了修复，并重新运行所有修改过的非机器学习实验脚本，将最新回测结果同步到本文档。'
)

# Fix list
doc.add_heading('修复内容', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '文件路径'
hdr_cells[1].text = '修复内容'
hdr_cells[2].text = '修复类型'

fixes = [
    ('src/run_optimization_experiments.py',
     '滚动夏普/波动率添加 .shift(1)；风险平价历史窗口改为 iloc[start_idx:date_idx] 排除当前日期',
     '计算方法'),
    ('src/run_week8_tutorial.py',
     '市值代理因子 rolling(12).std() 添加 .shift(1)',
     '计算方法'),
    ('src/run_experiments.py',
     '流动性过滤改为使用初始24个月数据计算统计量，避免全样本前视偏差',
     '静态筛选'),
    ('src/strategy/daily_tactician.py',
     'Z-score的 expanding().mean/std 添加 .shift(1)；阈值校准的 expanding().quantile() 添加 .shift(1)',
     '扩展窗口'),
    ('src/run_epu_weight_experiment.py',
     'EPU扩展窗口统计量及min/max归一化均添加 .shift(1)',
     '扩展窗口'),
    ('src/run_turnover_epu_experiment.py',
     'EPU阈值及中位数权重调整的 expanding() 统计添加 .shift(1)',
     '扩展窗口'),
]

for filepath, content, ftype in fixes:
    row_cells = table.add_row().cells
    row_cells[0].text = filepath
    row_cells[1].text = content
    row_cells[2].text = ftype

# Updated results
doc.add_heading('修正后实验结果汇总', level=2)

table2 = doc.add_table(rows=1, cols=5)
table2.style = 'Light Grid Accent 1'
hdr2 = table2.rows[0].cells
hdr2[0].text = '实验脚本'
hdr2[1].text = '关键指标'
hdr2[2].text = '修复前'
hdr2[3].text = '修复后'
hdr2[4].text = '变化'

results = [
    ('run_experiments.py', '基线年化换手', '1147.79%', '655.95%', '-491.84%'),
    ('run_optimization_experiments.py', '综合最优(年化/夏普/回撤)', '10.28%/0.74/-19.03%', '10.28%/0.74/-19.03%', '基本不变'),
    ('run_week8_tutorial.py', '全样本年化/夏普', '~19.23% / ~1.11', '19.88% / 1.04', '+0.65% / -0.07'),
    ('run_week8_tutorial.py', 'CH-3 Alpha (t值)', '14.13% (t=6.10)', '12.27% (t=5.10)', '归因更稳健'),
    ('run_turnover_epu_experiment.py', 'EPU高位降仓(p75)', '17.21% / 0.71 / -52.13%', '17.08% / 0.72 / -49.20%', '夏普微升，回撤改善'),
    ('run_turnover_epu_experiment.py', 'MA+EPU组合', '21.74% / 0.92 / -32.63%', '21.74% / 0.92 / -32.63%', '无变化'),
]

for script, metric, before, after, change in results:
    row = table2.add_row().cells
    row[0].text = script
    row[1].text = metric
    row[2].text = before
    row[3].text = after
    row[4].text = change

# Key findings
doc.add_heading('关键发现', level=2)
findings = [
    '基线实验换手率显著下降：修复后基线年化换手率从 1147.79% 降至 655.95%，原因是流动性过滤从全样本改为历史窗口计算，保留了更多股票，组合更稳定。',
    'Week 8 子样本稳健性：全样本年化收益从 ~19.23% 修正为 19.88%，CH-3 alpha 从 14.13% 修正为 12.27%（t=5.10），子样本差异更明显。',
    'EPU 实验：高位降仓阈值添加 .shift(1) 后，p75 方案夏普从 0.71 微升至 0.72，回撤从 -52.13% 改善至 -49.20%。',
    '优化实验：综合最优指标（10.28%/0.74/-19.03%）基本不变，说明前视偏差对核心结论影响有限。',
]
for finding in findings:
    doc.add_paragraph(finding, style='List Bullet')

# Disclaimer
doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('数据一致性声明：').bold = True
p.add_run('所有修改过的实验脚本（除机器学习部分外）均已重新运行，上述表格中的"修复后"列对应最新回测结果。未列出的实验（如 Week 7 ML 实验）未因前视偏差修复而重新运行。')

doc.save(doc_path)
print(f'已更新 docx: {doc_path}')
